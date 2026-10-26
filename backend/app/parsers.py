"""Resume text extraction for PDF and DOCX uploads."""
from __future__ import annotations

import io

import pdfplumber
from docx import Document

SUPPORTED_EXTENSIONS = {".pdf", ".docx"}


class UnsupportedFileType(Exception):
    pass


class EmptyResumeText(Exception):
    pass


def extract_text(filename: str, content: bytes) -> str:
    """Extract plain text from an uploaded PDF or DOCX resume.

    Raises UnsupportedFileType / EmptyResumeText on bad input so the
    API layer can turn them into clean 4xx responses.
    """
    lower = filename.lower()

    if lower.endswith(".pdf"):
        text = _extract_pdf(content)
    elif lower.endswith(".docx"):
        text = _extract_docx(content)
    else:
        raise UnsupportedFileType(
            f"Unsupported file type for '{filename}'. Please upload a .pdf or .docx file."
        )

    text = text.strip()
    if len(text) < 30:
        raise EmptyResumeText(
            "Could not extract readable text from this file. "
            "It may be a scanned/image-only PDF, password-protected, or empty."
        )
    return text


def _extract_pdf(content: bytes) -> str:
    chunks: list[str] = []
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            chunks.append(page_text)
    return "\n".join(chunks)


def _extract_docx(content: bytes) -> str:
    doc = Document(io.BytesIO(content))
    chunks: list[str] = []
    for para in doc.paragraphs:
        if para.text.strip():
            chunks.append(para.text)
    # also pull text out of tables, resumes sometimes use them for layout
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    chunks.append(cell.text)
    return "\n".join(chunks)
